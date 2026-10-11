import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def apply_callout_style(cell, accent_color="1E3A8A", bg_color="F8FAFC"):
    """
    Transforme une cellule de tableau en 'Callout Box' stylisée :
    - Bordure gauche épaisse colorée
    - Fond très léger
    - Padding généreux
    """
    tcPr = cell._tc.get_or_add_tcPr()
    
    # Fond
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
    tcPr.append(shd)
    
    # Bordure gauche uniquement (size 24 = 3pt)
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{accent_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    # Padding (Marges internes)
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="160" w:type="dxa"/>
            <w:bottom w:w="160" w:type="dxa"/>
            <w:left w:w="240" w:type="dxa"/>
            <w:right w:w="240" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_tufte_table(table):
    """
    Applique le style Edward Tufte à un tableau python-docx.
    Supprime toutes les bordures verticales, garde uniquement de fines horizontales.
    """
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideH w:val="single" w:sz="2" w:space="0" w:color="F1F5F9"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    # Replace existing borders if present, else append
    existing_borders = tblPr.xpath('w:tblBorders')
    if existing_borders:
        tblPr.remove(existing_borders[0])
    tblPr.append(borders)
    
    # Padding des cellules par défaut du tableau
    tblCellMar = parse_xml(f'''
        <w:tblCellMar {nsdecls("w")}>
            <w:top w:w="100" w:type="dxa"/>
            <w:bottom w:w="100" w:type="dxa"/>
            <w:left w:w="100" w:type="dxa"/>
            <w:right w:w="100" w:type="dxa"/>
        </w:tblCellMar>
    ''')
    existing_mar = tblPr.xpath('w:tblCellMar')
    if existing_mar:
        tblPr.remove(existing_mar[0])
    tblPr.append(tblCellMar)

def create_kpi_block(doc, metrics):
    """
    Crée un bloc de KPI géants sans bordure.
    metrics = [("42%", "CROISSANCE"), ("1.2M", "REVENUS")]
    """
    table = doc.add_table(rows=2, cols=len(metrics))
    # Supprimer toutes les bordures
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="none"/><w:bottom w:val="none"/>
            <w:left w:val="none"/><w:right w:val="none"/>
            <w:insideH w:val="none"/><w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    existing_borders = tblPr.xpath('w:tblBorders')
    if existing_borders:
        tblPr.remove(existing_borders[0])
    tblPr.append(borders)

    for idx, (val, label) in enumerate(metrics):
        # Value (Ligne 0)
        cell_val = table.cell(0, idx)
        p_val = cell_val.paragraphs[0]
        run_val = p_val.add_run(val)
        run_val.font.size = Pt(40)
        run_val.font.bold = True
        run_val.font.color.rgb = RGBColor(30, 58, 138) # Accent primaire
        
        # Label (Ligne 1)
        cell_lbl = table.cell(1, idx)
        p_lbl = cell_lbl.paragraphs[0]
        run_lbl = p_lbl.add_run(label.upper())
        run_lbl.font.size = Pt(9)
        run_lbl.font.color.rgb = RGBColor(100, 116, 139) # Muted
