#!/usr/bin/env python
"""Re-open a generated .pptx and either LINT it (default) or print an inventory.

Two modes
---------
* (default) ``lint``  -- run the Prznt Perfect design-system checks. Prints one
  ``slide N: ...`` line per violation and exits **1** if anything failed, **0**
  when clean. Calibrated so ``build/canonical.pptx`` passes with 0 violations.
* ``--inventory``     -- the old behaviour: per-slide shape inventory (type,
  position/size in pt, text excerpt, off-slide flag). Always exits 0.

The linter enforces the strictly-flat, sharp-cornered, Arial-only system:

  1. Shadows      -- a shape carries a drop/inner shadow, OR an auto-shape
                     inherits the theme shadow (effectRef) without the explicit
                     empty ``<a:effectLst/>`` suppression our primitives emit.
  2. Rounded      -- a FILLED ``prstGeom prst="roundRect"`` (filled shapes are
                     sharp; the unfilled outlined cover/closing pills are exempt).
  3. Fonts        -- any run that is not Arial, or is bold / italic.
  4. Type scale   -- any run size off the {110,72,64,48,36,32,26,24,20,16,14}px scale.
  5. Text colour  -- any run colour outside the 14 palette tokens.
  6. Badges       -- text baked INSIDE a filled shape (only tables may hold text
                     in a filled cell; cards are a fill-only rect + separate
                     transparent text boxes).
  7. Eyebrows     -- on a slide with a 64px title, any smaller text whose top is
                     above the title's top (cover 72px / title-less slides exempt).
  8. Pagination   -- counter-like text ("2/15", "Page ...") must be 14px and
                     un-filled; filled pagination plates are rejected.
  9. Theme colors -- the theme's clrScheme must carry the Prznt Perfect palette
                     (Background/Text/Accent slots per theme.COLOR_MAP), so the
                     PowerPoint color picker offers only design-system colors.

Run:  .venv/bin/python scripts/inspect_pptx.py [path] [--inventory]
"""

import json
import os
import re
import sys

from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn


# --------------------------------------------------------------------------- #
# inventory mode (unchanged behaviour)
# --------------------------------------------------------------------------- #
def _pt(emu):
    return round(Emu(emu).pt, 1)


def inspect(path):
    prs = Presentation(path)
    slide_w, slide_h = prs.slide_width, prs.slide_height
    print("=" * 78)
    print(f"FILE: {path}")
    print(f"valid: re-parsed OK | slides: {len(prs.slides)} | "
          f"slide size: {_pt(slide_w)} x {_pt(slide_h)} pt "
          f"({round(Emu(slide_w).inches, 3)} x {round(Emu(slide_h).inches, 3)} in)")
    print("=" * 78)

    total = 0
    off_slide = 0
    for i, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        total += len(shapes)
        print(f"\n-- Slide {i:02d}: {len(shapes)} shapes " + "-" * 40)
        for sh in shapes:
            kind = str(sh.shape_type).split()[0] if sh.shape_type else "?"
            L, T, W, H = _pt(sh.left), _pt(sh.top), _pt(sh.width), _pt(sh.height)
            excerpt = ""
            if sh.has_text_frame:
                txt = sh.text_frame.text.replace("\n", " / ").strip()
                if txt:
                    excerpt = f'  "{txt[:52]}"' + ("…" if len(txt) > 52 else "")
            flag = ""
            if (L < -0.5 or T < -0.5
                    or L + W > _pt(slide_w) + 0.5 or T + H > _pt(slide_h) + 0.5):
                flag = "  <<< OFF-SLIDE"
                off_slide += 1
            print(f"   {kind:11s} @({L:6.1f},{T:6.1f}) {W:6.1f}x{H:5.1f}pt"
                  f"{flag}{excerpt}")

    print("\n" + "=" * 78)
    print(f"TOTAL shapes: {total} | off-slide: {off_slide} | "
          f"status: {'OK' if off_slide == 0 else 'CHECK OFF-SLIDE SHAPES'}")
    print("=" * 78)
    return prs


# --------------------------------------------------------------------------- #
# linter
# --------------------------------------------------------------------------- #
# px type scale -> DrawingML sz (centipoints). px->pt is x0.5, sz = pt*100 = px*50.
_SCALE_PX = [110, 72, 64, 48, 36, 32, 26, 24, 20, 16, 14]  # 110px=55pt cover; 24px=12pt subhead
_ALLOWED_SZ = {px * 50 for px in _SCALE_PX}
_TITLE_SZ = 64 * 50            # 3200 -- the H1 title kegel
_PAGINATION_SZ = 14 * 50       # 700  -- the required pagination kegel

_TOKENS_PATH = os.path.normpath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "..",
    "references", "design-tokens.json"))


def _load_palette(path=_TOKENS_PATH):
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    return {spec["hex"].lstrip("#").upper() for spec in data["colors"].values()}


_PALETTE = _load_palette()

# Counter-like text that must be a bare 14px un-filled page number.
_PAGINATION_RE = [
    re.compile(r"\d+\s*/\s*\d+"),            # 2/15
    re.compile(r"\bpage\b", re.I),           # Page ...
    re.compile(r"\bslide\s*\d", re.I),       # Slide 2
    re.compile(r"\b\d+\s+of\s+\d+\b", re.I),  # 2 of 15
]


def _looks_like_pagination(txt):
    return any(rx.search(txt) for rx in _PAGINATION_RE)


def _find(el, tag):
    return None if el is None else el.find(qn(tag))


def _sp_text(sp):
    return "".join(t.text or "" for t in sp.iter(qn("a:t")))


def _run_sizes(sp):
    return [int(r.get("sz")) for r in sp.iter(qn("a:rPr")) if r.get("sz")]


def _sp_top(sp):
    off = _find(_find(_find(sp, "p:spPr"), "a:xfrm"), "a:off")
    return None if off is None else int(off.get("y"))


def _sp_box(sp):
    """(left, top, right) in EMU from the shape's xfrm, or None."""
    xfrm = _find(_find(sp, "p:spPr"), "a:xfrm")
    off, ext = _find(xfrm, "a:off"), _find(xfrm, "a:ext")
    if off is None or ext is None:
        return None
    x = int(off.get("x"))
    return x, int(off.get("y")), x + int(ext.get("cx"))


def _is_filled(sp):
    """True if the shape renders a fill (and could hide text like a badge)."""
    spPr = _find(sp, "p:spPr")
    if spPr is None:
        return False
    for tag in ("a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill"):
        if spPr.find(qn(tag)) is not None:
            return True
    if spPr.find(qn("a:noFill")) is not None:
        return False
    # No explicit fill node -> may inherit a theme fill via the style fillRef.
    fill_ref = _find(_find(sp, "p:style"), "a:fillRef")
    return fill_ref is not None and fill_ref.get("idx") not in (None, "0")


def _check_shadow(sp, i, out):
    spPr = _find(sp, "p:spPr")
    eff = _find(spPr, "a:effectLst")
    if eff is not None and (eff.find(qn("a:outerShdw")) is not None
                            or eff.find(qn("a:innerShdw")) is not None):
        out.append(f"slide {i}: shape carries a drop/inner shadow "
                   f"(effectLst outerShdw/innerShdw) -- the system is flat")
        return
    eff_ref = _find(_find(sp, "p:style"), "a:effectRef")
    if eff_ref is not None and eff_ref.get("idx") not in (None, "0") \
            and eff is None:
        out.append(f"slide {i}: shape inherits the theme shadow "
                   f"(effectRef idx={eff_ref.get('idx')}) without an explicit "
                   f"<a:effectLst/> -- build it via primitives.add_rect")


def _check_roundrect(spTree, i, out):
    # Filled shapes are always sharp. The ONE sanctioned rounded shape is the
    # cover / closing date+quarter pill: an UNFILLED, outlined chip (client rule
    # 2026-08-06). Flag only FILLED rounded rectangles.
    for sp in spTree.iter(qn("p:sp")):
        spPr = _find(sp, "p:spPr")
        g = _find(spPr, "a:prstGeom")
        if g is not None and g.get("prst") == "roundRect" and _is_filled(sp):
            out.append(f"slide {i}: filled rounded rectangle "
                       f"(prstGeom prst=roundRect) -- filled shapes are sharp; "
                       f"only the unfilled outlined cover/closing pills round")


def _check_runs(spTree, i, out):
    for rPr in spTree.iter(qn("a:rPr")):
        latin = rPr.find(qn("a:latin"))
        face = latin.get("typeface") if latin is not None else None
        if face is not None and face != "Arial":
            out.append(f"slide {i}: non-Arial font '{face}' "
                       f"-- Arial Regular is the only family")
        if rPr.get("b") in ("1", "true"):
            out.append(f"slide {i}: bold run -- the system has no bold weight")
        if rPr.get("i") in ("1", "true"):
            out.append(f"slide {i}: italic run -- the system has no italic")
        sz = rPr.get("sz")
        if sz is not None and int(sz) not in _ALLOWED_SZ:
            out.append(f"slide {i}: off-scale type size {int(sz) / 50:g}px "
                       f"-- allowed {{{','.join(map(str, _SCALE_PX))}}}px")
        sf = rPr.find(qn("a:solidFill"))
        srgb = sf.find(qn("a:srgbClr")) if sf is not None else None
        if srgb is not None:
            val = (srgb.get("val") or "").upper()
            if val and val not in _PALETTE:
                out.append(f"slide {i}: text colour #{val} is outside the "
                           f"14-token palette")


def _check_badges_eyebrows_pagination(spTree, i, out):
    entries = []          # (box, max_sz) per text shape, for the eyebrow rule
    title_boxes = []
    for sp in spTree.iter(qn("p:sp")):
        # -- shadow is per-shape (handled here to reuse the sp walk) --
        _check_shadow(sp, i, out)

        text = _sp_text(sp).strip()

        # (6) text inside a filled shape = badge (tables are graphicFrames, not sp)
        if text and _is_filled(sp):
            excerpt = text[:24] + ("…" if len(text) > 24 else "")
            out.append(f"slide {i}: text inside a filled shape (badge) -- "
                       f'"{excerpt}". Use a fill-only rect + separate text boxes')

        if not text:
            continue
        sizes = _run_sizes(sp)
        box = _sp_box(sp)
        if not sizes or box is None:
            continue
        max_sz = max(sizes)
        entries.append((box, max_sz))
        if max_sz == _TITLE_SZ:
            title_boxes.append(box)

        # (8) pagination
        if _looks_like_pagination(text):
            if max_sz != _PAGINATION_SZ:
                out.append(f"slide {i}: pagination text \"{text[:24]}\" is "
                           f"{max_sz / 50:g}px -- must be 14px")
            if _is_filled(sp):
                out.append(f"slide {i}: pagination \"{text[:24]}\" sits on a "
                           f"filled plate -- page numbers are un-filled text")

    # (7) eyebrow: smaller text above the top of a 64px title AND horizontally
    # overlapping it (a real kicker over the heading). Text off to the side that
    # merely sits higher -- e.g. the top-right logo/company name (R41) -- is NOT
    # an eyebrow, so require the x-ranges to overlap the title's.
    if title_boxes:
        title_top = min(b[1] for b in title_boxes)
        t_left = min(b[0] for b in title_boxes)
        t_right = max(b[2] for b in title_boxes)
        for (left, top, right), max_sz in entries:
            if (max_sz < _TITLE_SZ and top < title_top
                    and left < t_right and right > t_left):
                out.append(f"slide {i}: eyebrow label ({max_sz / 50:g}px) sits "
                           f"above the {_TITLE_SZ // 50}px title -- no labels "
                           f"over the heading")


def _check_theme(prs, out):
    """(9) The theme color scheme must equal the Prznt Perfect brand palette."""
    try:
        from przntperfect.theme import COLOR_MAP, THEME_NAME, read_clrScheme
        from przntperfect.tokens import TOKENS
    except Exception:
        out.append("theme: cannot import przntperfect.theme to verify the "
                   "brand color scheme")
        return
    expected = {slot: TOKENS.hex(name).lstrip("#").upper()
                for slot, name in COLOR_MAP}
    schemes = read_clrScheme(prs)
    if not schemes:
        out.append("theme: no clrScheme found -- the deck carries no theme "
                   "palette (expected the Prznt Perfect scheme)")
        return
    for partname, name, colors in schemes:
        if name != THEME_NAME:
            out.append(f"theme {partname}: scheme name is {name!r}, expected "
                       f"{THEME_NAME!r}")
        for slot, want in expected.items():
            got = (colors.get(slot) or "").upper()
            if got != want:
                out.append(f"theme {partname}: {slot} is #{got or 'none'}, "
                           f"expected #{want} -- picker palette is off-brand")


def lint(path):
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        spTree = slide.shapes._spTree
        _check_roundrect(spTree, i, out)
        _check_runs(spTree, i, out)
        _check_badges_eyebrows_pagination(spTree, i, out)
    _check_theme(prs, out)

    print("=" * 78)
    print(f"LINT: {path}")
    print("=" * 78)
    for line in out:
        print(line)
    print("-" * 78)
    print(f"{len(out)} violation{'' if len(out) == 1 else 's'}"
          f" | status: {'PASS' if not out else 'FAIL'}")
    print("=" * 78)
    return len(out)


# --------------------------------------------------------------------------- #
if __name__ == "__main__":
    _default = os.path.normpath(
        os.path.join(os.path.dirname(os.path.abspath(__file__)),
                     "..", "build", "canonical.pptx")
    )
    _args = sys.argv[1:]
    _inventory = "--inventory" in _args
    _paths = [a for a in _args if not a.startswith("--")]
    _path = _paths[0] if _paths else _default

    if _inventory:
        inspect(_path)
        sys.exit(0)
    sys.exit(1 if lint(_path) else 0)
