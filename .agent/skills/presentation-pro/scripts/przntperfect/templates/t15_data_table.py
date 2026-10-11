"""15 - Data table (Figma node 562:99), rendered as a NATIVE PowerPoint table.

Client-approved (2026-07-23) exception to the "text + lines" rule, mirroring the
agenda (slide 02): the data table is a real ``GraphicFrame`` table object so rows
can be added/removed and text stays freely editable. Only HORIZONTAL separators
are drawn (no verticals, no fills, default table style disabled). The geometry
reproduces the previous text+hairline render exactly (that render matched the
reference):

* Header underline (below the "Column 1" row) is HEAVIER -- 1.5pt (=3px, w=19050),
  centered on the row boundary at y392.5 so it covers the reference's 3 solid
  pixel rows 391-393. Realised as the header row's bottom border.
* Six data-row underlines are 0.5pt (=1px, w=6350), each realised as that data
  row's bottom border, landing on the measured reference rows 493/594/695/795/
  896/991. The last data row is shorter than the others (canonical).
* All separators are Gray-Dark #A8B1C7. No top border (nothing above the header),
  no left/right borders anywhere -> only the seven horizontal lines show.
* The half-pixel trick (t02): row boundaries sit on nnn.5 px so PowerPoint's
  CENTER-aligned border stroke lands on a single pixel row (1px lines) / the
  symmetric 391-393 band (the 3px header line).

Cell text is styled exactly like primitives.add_text (Arial Regular, exact
spcPts line pitch, -1% tracking): header 32px (spc -16), data 26px (spc -13),
all Graphite #454E63. Cells use marT=0 + anchor="ctr": text is vertically
centered in the row (client rule 2026-07-26). Column 2's header sits 16px
right of its data (reference: header x434 vs data x418) via an extra marL on the
header col-2 cell over a shared column edge at x418.

The "Ultra-high-resolution / sensing" H2 heading and the footer are unchanged.
"""

from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

from .. import primitives as P
from ._common import white_bg, light_footer, GRAPHITE, GRAY_DARK
from ..units import PXt, tracking_spc

# --- content ---------------------------------------------------------------- #
HEADER_TEXT = "Column 1"                                     # x5 across the header
VALUES = ["0.18 msec", "23.43 msec", "12.43 msec", "3.43 msec", "0.18 msec"]
NROWS = 7                                                    # 1 header + 6 data
NCOLS = 5

# --- table geometry (px on the 1920x1080 canvas) ---------------------------- #
TABLE_X = 50
# Frame top. No border renders here (header lnT = noFill), so this only sets the
# header row's height and where its text anchors; it need not be on the .5 grid.
TABLE_Y = 320.5
TABLE_W = 1820
# Column widths -> left edges 50 / 418 / 786 / 1154 / 1522, right edge 1870.
# Column 2's edge is the DATA origin x418; the header value is nudged right via
# marL (see COL2_HEADER_MARL). Sum = 1820.
COL_WIDTHS = [368, 368, 368, 368, 348]

# Row lower boundaries (px), all on the nnn.5 grid so a center-stroked border
# lands cleanly: the 1.5pt header line at 392.5 covers rows 391-393; each 0.5pt
# data line lands on the single measured reference row (493/594/695/795/896/991).
# Row tops are TABLE_Y then each previous boundary. Last data row (896.5->991.5,
# 95px) is shorter than the ~101px rows above it -- canonical.
ROW_BOTTOMS = [392.5, 493.5, 594.5, 695.5, 795.5, 896.5, 991.5]

# Border weights (EMU) and colour.
HEADER_LNB = 19050          # 1.5pt = 3px  (heavy header underline)
DATA_LNB = 6350             # 0.5pt = 1px  (data-row underlines)
SEP_HEX = str(GRAY_DARK)    # "A8B1C7"

# Per-row font size. Text is vertically centered in the row (client rule
# 2026-07-26): marT=0 + anchor="ctr", so no per-row glyph-top calibration.
SIZES = [32, 26, 26, 26, 26, 26, 26]

# Header column 2 only: data edge at x418, header value at x434 (+16px).
COL2_HEADER_MARL = 16

# "No Style, No Grid" table style -- no header fill, no banding, no gridlines.
NO_STYLE_NO_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _disable_table_style(table):
    """Kill the default tableStyle (header/banding fills) so cells stay clean."""
    table.first_row = False
    table.last_row = False
    table.first_col = False
    table.last_col = False
    table.horz_banding = False
    table.vert_banding = False
    tblPr = table._tbl.tblPr
    styleId = tblPr.find(qn("a:tableStyleId"))
    if styleId is None:
        styleId = parse_xml(
            '<a:tableStyleId xmlns:a="http://schemas.openxmlformats.org/'
            f'drawingml/2006/main">{NO_STYLE_NO_GRID}</a:tableStyleId>'
        )
        tblPr.append(styleId)
    else:
        styleId.text = NO_STYLE_NO_GRID


def _fill_cell_text(cell, text, size_px, color):
    """Style a cell's text exactly like primitives.add_text (Arial, exact spc)."""
    tf = cell.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    # Exact line pitch in points (spcPts), never a float multiple. See units.
    para.line_spacing = PXt(size_px)                 # size_px * 100%
    run = para.add_run()
    run.text = text
    f = run.font
    f.name = "Arial"
    f.size = PXt(size_px)
    f.bold = False
    f.italic = False
    f.color.rgb = color
    run._r.get_or_add_rPr().set("spc", str(tracking_spc(size_px)))  # -16 / -13


def _ln(tag, w_emu):
    """A cell-border XML fragment: a Gray-Dark line of weight ``w_emu``, or,
    when ``w_emu`` is 0, an explicit no-border."""
    if w_emu:
        return (
            f'<a:{tag} w="{w_emu}" cap="flat" cmpd="sng" algn="ctr">'
            f'<a:solidFill><a:srgbClr val="{SEP_HEX}"/></a:solidFill>'
            f'<a:prstDash val="solid"/></a:{tag}>'
        )
    return f'<a:{tag}><a:noFill/></a:{tag}>'


def _set_cell_box(cell, marL_px, lnB_w):
    """Rewrite the cell's tcPr: zeroed insets (bar the calibration marL),
    center-anchor, borders, noFill.

    marT=0 + anchor="ctr": text is vertically centered in the row
    (client rule 2026-07-26). Only the bottom border is drawn (each row draws
    its own underline); lnT is an explicit no-border so the shared boundary is
    stroked exactly once by the row ABOVE (the header's 1.5pt / a data row's
    0.5pt). lnL/lnR are no-borders and the fill is noFill so the white slide
    shows through.
    """
    tc = cell._tc
    existing = tc.find(qn("a:tcPr"))
    if existing is not None:
        tc.remove(existing)
    marL = int(PXt(marL_px))          # EMU
    xml = (
        '<a:tcPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'marL="{marL}" marR="0" marT="0" marB="0" anchor="ctr">'
        f'{_ln("lnL", 0)}{_ln("lnR", 0)}'
        f'{_ln("lnT", 0)}{_ln("lnB", lnB_w)}'
        "<a:noFill/>"
        "</a:tcPr>"
    )
    tc.append(parse_xml(xml))          # tcPr follows txBody in CT_TableCell


def build(deck):
    s = deck.add_slide()
    white_bg(s)

    # H2 heading -- two single-line boxes (unchanged; see history in git).
    P.add_text(s, 50, 44, 800, 90, "Ultra-high-resolution",
               size_px=64, line_height_pct=100, color=GRAPHITE)
    P.add_text(s, 50, 98, 800, 90, "sensing",
               size_px=64, line_height_pct=100, color=GRAPHITE)

    # Derive row tops/heights from the boundaries (geometry unchanged).
    row_tops = [TABLE_Y] + ROW_BOTTOMS[:-1]
    row_heights = [ROW_BOTTOMS[i] - row_tops[i] for i in range(NROWS)]

    gf = s.shapes.add_table(
        NROWS, NCOLS,
        PXt(TABLE_X), PXt(TABLE_Y), PXt(TABLE_W), PXt(sum(row_heights)),
    )
    table = gf.table
    _disable_table_style(table)
    for c in range(NCOLS):
        table.columns[c].width = PXt(COL_WIDTHS[c])

    for r in range(NROWS):
        table.rows[r].height = PXt(row_heights[r])
        is_header = (r == 0)
        size = SIZES[r]
        lnB = HEADER_LNB if is_header else DATA_LNB
        for c in range(NCOLS):
            cell = table.cell(r, c)
            text = HEADER_TEXT if is_header else VALUES[c]
            _fill_cell_text(cell, text, size, GRAPHITE)
            marL = COL2_HEADER_MARL if (is_header and c == 1) else 0
            _set_cell_box(cell, marL, lnB)

    light_footer(s)
    return s
