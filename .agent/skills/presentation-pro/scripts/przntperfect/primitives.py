"""Drawing primitives for the Prznt Perfect Slides engine.

Every primitive works on a python-pptx slide and positions/sizes in px (converted
to pt via units.PXt). The system is strictly flat (R12): shapes get no shadow and
no decorative outline unless explicitly requested.

Text model
----------
`add_text` accepts either a plain string or a list of ``(text, RGBColor)``
segments (for two-tone headings). Any ``"\n"`` inside a segment's text becomes an
explicit DrawingML line break (``<a:br/>``) -- never a literal U+2028. Every run
carries Arial Regular, the size, and -1% tracking (rPr@spc).
"""

from math import ceil

from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

from .units import PXt, tracking_spc
from .textbalance import (normalize_dashes, glue_short_words, glue_all_words,
                          text_width_px, line_count)
from .tokens import TOKENS

_ACCENT_RGB = TOKENS.color("High")   # the accent colour (#00A1FF) for R37

_ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
_ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


# --------------------------------------------------------------------------- #
# low-level shape helpers
# --------------------------------------------------------------------------- #
def _flatten(shape):
    """Strictly-flat look: no inherited shadow."""
    shape.shadow.inherit = False


def _no_line(shape):
    shape.line.fill.background()


def _fill(shape, rgb):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb


def add_rect(slide, x_px, y_px, w_px, h_px, rgb, shape_type=MSO_SHAPE.RECTANGLE):
    """A flat, borderless, filled rectangle (used for backgrounds, bars, lines)."""
    shp = slide.shapes.add_shape(
        shape_type, PXt(x_px), PXt(y_px), PXt(w_px), PXt(h_px)
    )
    _fill(shp, rgb)
    _no_line(shp)
    _flatten(shp)
    return shp


# --------------------------------------------------------------------------- #
# backgrounds
# --------------------------------------------------------------------------- #
_BLEED_PX = 4  # off-canvas overhang for the gradient background (see below)


def add_solid_background(slide, rgb):
    """Solid slide background via the REAL background property (p:bg).

    Not a full-bleed rectangle: a rect is selectable/movable/deletable by the
    end user, while the true background is inert and shows up properly in
    PowerPoint's Format Background. (The cover gradient still uses a bled
    rectangle - see add_vertical_gradient_background - because the gradFill
    background rendering leaves a seam there.)
    """
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb
    return None


def add_vertical_gradient_background(slide, stops):
    """Full-bleed rectangle with a vertical (top->bottom) linear gradient.

    ``stops`` is a list of ``(pos_thousandths_pct, 'RRGGBB')`` tuples, e.g.
    ``[(0, '454E63'), (100000, '8A93A9')]``.
    python-pptx has no multi-stop gradient API, so the gradFill is built as XML.

    The rectangle is bled ``_BLEED_PX`` past every canvas edge. A solid-fill
    full-bleed rect covers row 0 cleanly, but the renderer leaves a 1 px
    top-left seam (bare white slide showing through) when the same preset
    rectangle carries this linear gradFill. Pushing the shape edge off-canvas
    removes the seam; the vertical colour mapping shifts < 0.5/255 (negligible).
    """
    b = _BLEED_PX
    shp = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, PXt(-b), PXt(-b), PXt(1920 + 2 * b), PXt(1080 + 2 * b)
    )
    _no_line(shp)
    _flatten(shp)

    sp_pr = shp._element.spPr
    # remove any existing fill node
    for tag in (
        "a:noFill", "a:solidFill", "a:gradFill", "a:blipFill",
        "a:pattFill", "a:grpFill",
    ):
        for el in sp_pr.findall(qn(tag)):
            sp_pr.remove(el)

    gs_xml = "".join(
        f'<a:gs pos="{pos}"><a:srgbClr val="{hexv}"/></a:gs>' for pos, hexv in stops
    )
    grad_xml = (
        '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f"<a:gsLst>{gs_xml}</a:gsLst>"
        '<a:lin ang="5400000" scaled="1"/>'  # 5400000 = 90deg -> top-to-bottom
        "</a:gradFill>"
    )
    grad = parse_xml(grad_xml)
    ln = sp_pr.find(qn("a:ln"))
    if ln is not None:
        ln.addprevious(grad)
    else:
        sp_pr.append(grad)
    return shp


# --------------------------------------------------------------------------- #
# text
# --------------------------------------------------------------------------- #
def _style_run(run, size_px, rgb, tracking=True, track_pt=None):
    f = run.font
    f.name = "Arial"
    f.size = PXt(size_px)          # px -> pt Length
    f.bold = False
    f.italic = False
    if rgb is not None:
        f.color.rgb = rgb
    # ``track_pt`` sets an ABSOLUTE letter spacing in points (PowerPoint's
    # "Condensed/Expanded by N pt"); e.g. -1.5 -> spc=-150. Overrides the default
    # -1% (size-relative) tracking. spc is DrawingML 1/100 pt.
    if track_pt is not None:
        run._r.get_or_add_rPr().set("spc", str(round(track_pt * 100)))
    elif tracking:
        run._r.get_or_add_rPr().set("spc", str(tracking_spc(size_px)))


def _emit_paragraph(paragraph, segments, size_px, tracking, track_pt=None):
    """Build runs (and <a:br/> breaks) for a paragraph from colored segments.

    ``segments`` = list of ``(text, rgb)``. A ``"\n"`` in a segment's text
    produces an explicit break placed immediately before the following run.
    """
    # In a TWO-TONE heading, the accent (High) run is a phrase that must not split
    # across lines leaving one word hanging (R37) -- glue the whole accent run.
    two_tone = len(segments) > 1
    pending_breaks = 0
    for text, rgb in segments:
        norm = normalize_dashes(text)
        if two_tone and rgb is not None and rgb == _ACCENT_RGB:
            parts = glue_all_words(norm).split("\n")     # accent phrase stays whole
        else:
            # glue short words so none ends a line (R33)
            parts = glue_short_words(norm).split("\n")
        for i, part in enumerate(parts):
            if i > 0:
                pending_breaks += 1
            if part == "":
                continue
            run = paragraph.add_run()
            run.text = part
            _style_run(run, size_px, rgb, tracking, track_pt)
            for _ in range(pending_breaks):
                run._r.addprevious(parse_xml(
                    '<a:br xmlns:a="http://schemas.openxmlformats.org/'
                    'drawingml/2006/main"/>'
                ))
            pending_breaks = 0


def add_text(
    slide, x_px, y_px, w_px, h_px, content, size_px,
    line_height_pct=100, align="l", anchor="t", color=None,
    tracking=True, wrap=True, left_inset_px=0, hug=True, track_pt=None,
):
    """Place a text box.

    ``content`` may be a ``str`` (uses ``color``) or a list of ``(text, rgb)``
    segments for two-tone / multi-color runs. Box margins are zeroed so px
    coordinates are exact (optional ``left_inset_px`` for e.g. chart labels).

    ``hug`` (default True) shrinks the box to the text like Figma HUG, so a
    selected box is snug with no dead space (R36). Unlike PowerPoint auto-fit
    (which only trims height, and only once the text is edited), the box is set to
    the MEASURED text bbox up front — width = the longest rendered line, height =
    line count x pitch — so it is tight the moment the deck opens, with no
    first-line "sag". Only TOP-anchored boxes hug (the anchor stays put). Width
    only shrinks for LEFT-aligned text (shrinking a centre/right box would move
    the copy); height always hugs. ``track_pt`` sets an absolute letter spacing in
    points (overrides the -1% default; used by the 110px cover heading).
    """
    if hug and anchor == "t":
        full = content if isinstance(content, str) else "".join(t for t, _ in content)
        hard = normalize_dashes(full).split("\n")
        n_render = sum(line_count(ln, w_px, size_px) for ln in hard) or 1
        pitch = size_px * line_height_pct / 100.0
        h_px = ceil(n_render * pitch) + 2
        if align == "l":                       # width-hug is safe only left-aligned
            # Measure UNTRACKED (widest) so a tracked/kerned render never clips.
            max_w = max((text_width_px(ln, size_px, tracking=False) for ln in hard),
                        default=0)
            w_px = min(w_px, ceil(max_w) + left_inset_px + 2)

    box = slide.shapes.add_textbox(PXt(x_px), PXt(y_px), PXt(w_px), PXt(h_px))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None                        # box is already the exact size
    tf.margin_left = PXt(left_inset_px)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    tf.vertical_anchor = _ANCHOR[anchor]

    para = tf.paragraphs[0]
    para.alignment = _ALIGN[align]
    # Exact line pitch in points (spcPts), Figma semantics: pitch = size x lh%.
    # NEVER assign a float multiple here (spcPct): renderers multiply it by the
    # font's own natural leading (~1.22em for Arial), inflating every multi-line
    # block by ~22% and sagging first lines by the extra half-leading.
    para.line_spacing = PXt(size_px * line_height_pct / 100.0)

    if isinstance(content, str):
        segments = [(content, color)]
    else:
        segments = content
    _emit_paragraph(para, segments, size_px, tracking, track_pt)
    return box


def add_two_tone(slide, x_px, y_px, w_px, h_px, segments, size_px, **kwargs):
    """Convenience alias: a heading made of several differently-colored runs."""
    return add_text(slide, x_px, y_px, w_px, h_px, segments, size_px, **kwargs)


# --------------------------------------------------------------------------- #
# image placeholder / hairline
# --------------------------------------------------------------------------- #
def add_image_placeholder(slide, x_px, y_px, w_px, h_px, rgb, radius_px=0):
    """White-Cool image placeholder.

    Reference renders show SHARP corners (strict pixel check: 0 px inset on all
    placeholders), so the default radius is 0. A positive ``radius_px`` switches
    to a rounded-rectangle whose corner radius is ``radius_px`` px.
    """
    if radius_px and radius_px > 0:
        shp = add_rect(
            slide, x_px, y_px, w_px, h_px, rgb,
            shape_type=MSO_SHAPE.ROUNDED_RECTANGLE,
        )
        # adjustment value is radius / min(w, h)
        try:
            shp.adjustments[0] = float(radius_px) / min(w_px, h_px)
        except Exception:
            pass
        return shp
    return add_rect(slide, x_px, y_px, w_px, h_px, rgb)


def add_hairline(slide, x_px, y_px, w_px, rgb, weight_px=1):
    """A thin horizontal separator drawn as a flat filled rectangle."""
    return add_rect(slide, x_px, y_px, w_px, weight_px, rgb)


# --------------------------------------------------------------------------- #
# picture (embedded image) -- e.g. the user's logo
# --------------------------------------------------------------------------- #
def add_picture(slide, image_path, x_px, y_px, w_px, h_px, align="r", valign="t"):
    """Embed an image, scaled to FIT WITHIN the ``w_px`` x ``h_px`` box while
    preserving its aspect ratio, then aligned inside the box (``align`` l/c/r,
    ``valign`` t/m/b). Used for a user-supplied logo (the one sanctioned raster in
    the deck — a brand asset, not a content image). Flat: no shadow."""
    from pptx.util import Emu
    pic = slide.shapes.add_picture(image_path, PXt(x_px), PXt(y_px))
    nw = Emu(pic.width).pt * 2.0                      # native size in canvas px
    nh = Emu(pic.height).pt * 2.0
    scale = min(w_px / nw, h_px / nh) if nw and nh else 1.0
    fw, fh = nw * scale, nh * scale
    pic.width, pic.height = PXt(fw), PXt(fh)
    lx = x_px + (w_px - fw) if align == "r" else (
        x_px + (w_px - fw) / 2 if align == "c" else x_px)
    ty = y_px + (h_px - fh) if valign == "b" else (
        y_px + (h_px - fh) / 2 if valign == "m" else y_px)
    pic.left, pic.top = PXt(lx), PXt(ty)
    try:
        pic.shadow.inherit = False
    except Exception:
        pass
    return pic


# --------------------------------------------------------------------------- #
# pill (rounded, outlined, unfilled text chip) -- cover / closing only
# --------------------------------------------------------------------------- #
def add_pill(slide, x_px, y_px, w_px, h_px, text, text_color, line_color,
             line_weight_px=1, size_px=20, line_height_pct=124,
             inset_px=20, tracking=True):
    """A fully-rounded, OUTLINED, UNFILLED chip with centered text.

    The ONE sanctioned rounded shape in the system (client rule 2026-08-06): the
    cover / closing date + quarter chips. It is a rounded-rectangle textbox with a
    1px outline (``line_color``), NO fill (the background shows through, so the
    linter's "badge" rule -- text inside a *filled* shape -- does not apply), and
    max corner rounding (adjustment 0.5 -> a true pill). Text is Arial with -1%
    tracking, centered horizontally AND vertically, with ``inset_px`` left/right
    insets. Flat (no shadow).
    """
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, PXt(x_px), PXt(y_px), PXt(w_px), PXt(h_px))
    shp.fill.background()                       # unfilled: NOT a badge
    shp.line.color.rgb = line_color
    shp.line.width = PXt(line_weight_px)        # 1px = 0.5pt hairline outline
    _flatten(shp)
    try:
        shp.adjustments[0] = 0.5                # max rounding -> a pill
    except Exception:
        pass

    tf = shp.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    tf.margin_left = PXt(inset_px)
    tf.margin_right = PXt(inset_px)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE      # center vertically
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER            # center horizontally
    para.line_spacing = PXt(size_px * line_height_pct / 100.0)
    run = para.add_run()
    run.text = text
    _style_run(run, size_px, text_color, tracking)
    return shp


# --------------------------------------------------------------------------- #
# circular arrow button (cover element) -- native flat shapes, no SVG import
# --------------------------------------------------------------------------- #
def add_arrow_button(slide, x_px, y_px, d_px, color, weight_px=1):
    """A circular 'next' button: an unfilled circle outline with a right-arrow
    inside, both in ``color`` (Figma node 670-1025, cover element).

    Built from NATIVE flat shapes — an OVAL outline plus three straight strokes
    (shaft + two chevron arms) — not an imported SVG/raster, so it is editable,
    theme-colored, and pixel-crisp at any zoom (the right way to add a simple
    vector mark; reserve SVG-file import for complex illustrations). The circle
    is unfilled, so it is not a badge and joins the pills as a sanctioned round
    outline (R12 exception)."""
    r = d_px / 2.0
    cx, cy = x_px + r, y_px + r
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, PXt(x_px), PXt(y_px),
                                    PXt(d_px), PXt(d_px))
    circle.fill.background()
    circle.line.color.rgb = color
    circle.line.width = PXt(weight_px)
    _flatten(circle)
    # arrow geometry (ratios measured from the Figma vector): the point sits at
    # +0.48r on the axis, the chevron arms reach +/-0.50r vertically.
    point = (cx + 0.48 * r, cy)
    tail = (cx - 0.48 * r, cy)
    top = (cx - 0.02 * r, cy - 0.50 * r)
    bot = (cx - 0.02 * r, cy + 0.50 * r)
    for (x1, y1), (x2, y2) in ((tail, point), (point, top), (point, bot)):
        ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                        PXt(x1), PXt(y1), PXt(x2), PXt(y2))
        ln.line.color.rgb = color
        ln.line.width = PXt(weight_px)
        try:
            ln.shadow.inherit = False
        except Exception:
            pass
    return circle


# --------------------------------------------------------------------------- #
# footer
# --------------------------------------------------------------------------- #
def add_footer(
    slide, caption_color, page=None, page_color=None,
    caption="Caption", cover=False,
):
    """Footer caption (bottom-left) and optional pagination (bottom-right).

    Cover (01): caption at (50,1016), no pagination. Slides 04-15: caption at
    (52,1033) and page number right-aligned to the content edge.

    Footer baselines are NOT co-linear in the reference: the caption glyphs top at
    y1033 (box 1033) while the pagination sits ~17px lower, glyph top y1050
    (box 1048). The old code stacked both boxes at y1030, so the page number
    rendered ~18px too high -- it landed clear ABOVE the reference "10", producing
    the doubled page number visible on every footer slide's mask. Measured on
    slides 04/06/12/13/15: REF page-num bbox yt1050 vs REN yt1032.
    """
    cx, cy = (50, 1016) if cover else (52, 1033)
    add_text(
        slide, cx, cy, 400, 30, caption,
        size_px=16, line_height_pct=100, align="l", color=caption_color,
    )
    if page is not None:
        # right edge at ~1867 (box 1667..1867, right-aligned): REF page-num glyph
        # right measured x1866 vs REN x1869 at the old box-right 1870. y1048 lands
        # the 14px glyph top at y1050 to match REF (was 1030 -> glyph 1032).
        add_text(
            slide, 1667, 1048, 200, 26, page,
            size_px=14, line_height_pct=110, align="r",
            color=page_color if page_color is not None else caption_color,
        )
