"""Reusable native-table helper: the Prznt Perfect "No-Style-No-Grid,
horizontal-borders-only" table.

The canonical templates 02/14/15 each inline this same construction (disable the
default table style, zero every cell inset, draw only horizontal borders in
Gray-Dark). Those are pixel-tuned and left untouched; this module packages the
identical technique so NEW components (the numbered-steps card and the metric
card) build real, row-addable ``GraphicFrame`` tables instead of loose shapes —
exactly what the client asked for ("сделать как таблица").

A cell holds one or more stacked paragraphs, each with its own size/color/pitch
(so a row can carry a big ``Num-Head`` number above ``Body`` text). Only
horizontal rules are drawn: an optional top rule on row 0 plus a per-row bottom
rule. Vertical borders and fills are always off, so whatever is behind the table
(the white slide or a White-Cool card) shows through.
"""

from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

from .units import PXt, tracking_spc
from .tokens import TOKENS

_GREY_LIGHT_HEX = TOKENS.hex("Grey-Light").lstrip("#").upper()   # separators
_GRAY_DARK_HEX = TOKENS.hex("Gray-Dark").lstrip("#").upper()     # chart-table rule

# "No Style, No Grid" built-in table style: no header fill, no banding, no grid.
NO_STYLE_NO_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

# Rule weights (EMU): 0.5pt = 1px (6350), 1.5pt = 3px (19050).
RULE_THIN = 6350
RULE_THICK = 19050


def _disable_table_style(table):
    table.first_row = table.last_row = False
    table.first_col = table.last_col = False
    table.horz_banding = table.vert_banding = False
    tblPr = table._tbl.tblPr
    styleId = tblPr.find(qn("a:tableStyleId"))
    if styleId is None:
        tblPr.append(parse_xml(
            '<a:tableStyleId xmlns:a="http://schemas.openxmlformats.org/'
            f'drawingml/2006/main">{NO_STYLE_NO_GRID}</a:tableStyleId>'))
    else:
        styleId.text = NO_STYLE_NO_GRID


def _ln(tag, w_emu, hexv):
    """Border fragment: a solid line of ``w_emu`` EMU, or an explicit no-border."""
    if w_emu:
        return (
            f'<a:{tag} w="{w_emu}" cap="flat" cmpd="sng" algn="ctr">'
            f'<a:solidFill><a:srgbClr val="{hexv}"/></a:solidFill>'
            f'<a:prstDash val="solid"/></a:{tag}>'
        )
    return f'<a:{tag}><a:noFill/></a:{tag}>'


def _set_cell_box(cell, marL_px, marT_px, top_w, bottom_w, hexv, anchor):
    tc = cell._tc
    old = tc.find(qn("a:tcPr"))
    if old is not None:
        tc.remove(old)
    marL = int(PXt(marL_px))
    marT = int(PXt(marT_px))
    xml = (
        '<a:tcPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'marL="{marL}" marR="0" marT="{marT}" marB="0" anchor="{anchor}">'
        f'{_ln("lnL", 0, hexv)}{_ln("lnR", 0, hexv)}'
        f'{_ln("lnT", top_w, hexv)}{_ln("lnB", bottom_w, hexv)}'
        "<a:noFill/>"
        "</a:tcPr>"
    )
    tc.append(parse_xml(xml))


def _fill_cell_paragraphs(cell, paras):
    """Write stacked paragraphs into a cell.

    ``paras`` = list of dicts: ``text``, ``size_px``, ``color`` (RGBColor),
    ``lh_pct`` (line-height %, default 100), ``space_before_px`` (default 0).
    Each run is Arial Regular with exact ``spcPts`` pitch and −1% tracking,
    identical to :func:`primitives.add_text`.
    """
    tf = cell.text_frame
    tf.word_wrap = True
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.LEFT
        size_px = p["size_px"]
        lh = p.get("lh_pct", 100)
        para.line_spacing = PXt(size_px * lh / 100.0)
        sb = p.get("space_before_px", 0)
        if sb:
            para.space_before = PXt(sb)
        run = para.add_run()
        run.text = p["text"]
        f = run.font
        f.name = "Arial"
        f.size = PXt(size_px)
        f.bold = False
        f.italic = False
        if p.get("color") is not None:
            f.color.rgb = p["color"]
        run._r.get_or_add_rPr().set("spc", str(tracking_spc(size_px)))


def stacked_rule_table(slide, x_px, y_px, w_px, rows, rule_hex=_GREY_LIGHT_HEX,
                       rule_weight_emu=RULE_THIN, top_rule=False, anchor="ctr",
                       marL_px=0):
    """Build a single-column horizontal-rule table and return its GraphicFrame.

    ``rows`` = list of dicts:
        ``paras``       -- stacked paragraphs for the cell (see
                           :func:`_fill_cell_paragraphs`).
        ``height_px``   -- row height.
        ``bottom_rule`` -- draw this row's bottom border (default True).
        ``marT_px``     -- extra top inset for this cell's text (default 0).
    ``top_rule`` draws a border above row 0. All rules are ``rule_hex`` at
    ``rule_weight_emu``; only horizontal rules are ever drawn.
    """
    n = len(rows)
    total_h = sum(r["height_px"] for r in rows)
    gf = slide.shapes.add_table(n, 1, PXt(x_px), PXt(y_px), PXt(w_px), PXt(total_h))
    table = gf.table
    _disable_table_style(table)
    table.columns[0].width = PXt(w_px)
    for r, row in enumerate(rows):
        table.rows[r].height = PXt(row["height_px"])
        cell = table.cell(r, 0)
        _fill_cell_paragraphs(cell, row["paras"])
        top_w = rule_weight_emu if (r == 0 and top_rule) else 0
        bot_w = rule_weight_emu if row.get("bottom_rule", True) else 0
        _set_cell_box(cell, marL_px, row.get("marT_px", 0),
                      top_w, bot_w, rule_hex, anchor)
    return gf


def segmented_underline_table(slide, x_px, y_px, content_widths, spacer_widths,
                              rows, rule_hex=_GRAY_DARK_HEX,
                              rule_weight_emu=RULE_THIN, anchor="ctr"):
    """Multi-column table whose row underline BREAKS across the spacer columns.

    Content columns are interleaved with narrow spacer columns
    (``c0 s0 c1 s1 … c(n-1)``). Only content cells carry a bottom border, so each
    row's underline shows as one SEGMENT under every content column with a GAP over
    every spacer — the technique the canonical chart table (t14) uses to line the
    underline up with the bars. ``content_widths`` (n) and ``spacer_widths`` (n-1)
    define the columns in px.

    ``rows`` = list of dicts:
        ``cells``       -- list of n cell specs (content columns only). Each spec
                           is a dict ``{text, size_px, color, lh_pct(=100),
                           marL_px(=0)}``.
        ``height_px``   -- row height.
        ``bottom_rule`` -- draw this row's (segmented) bottom border (default True).
        ``marT_px``     -- extra top inset (default 0).
    """
    n = len(content_widths)
    if len(spacer_widths) != n - 1:
        raise ValueError("segmented_underline_table: need n-1 spacer widths")

    col_widths = []
    for i in range(n):
        col_widths.append(content_widths[i])
        if i < n - 1:
            col_widths.append(spacer_widths[i])
    ncols = len(col_widths)
    nrows = len(rows)
    total_h = sum(r["height_px"] for r in rows)

    gf = slide.shapes.add_table(nrows, ncols, PXt(x_px), PXt(y_px),
                                PXt(sum(col_widths)), PXt(total_h))
    table = gf.table
    _disable_table_style(table)
    for c, w in enumerate(col_widths):
        table.columns[c].width = PXt(w)

    for r, row in enumerate(rows):
        table.rows[r].height = PXt(row["height_px"])
        bot = rule_weight_emu if row.get("bottom_rule", True) else 0
        marT = row.get("marT_px", 0)
        for ci in range(ncols):
            cell = table.cell(r, ci)
            if ci % 2 == 0:                          # content column
                spec = row["cells"][ci // 2]
                _fill_cell_paragraphs(cell, [spec])
                _set_cell_box(cell, spec.get("marL_px", 0), marT,
                              0, bot, rule_hex, anchor)
            else:                                    # spacer column: no border
                _fill_cell_paragraphs(cell, [])
                _set_cell_box(cell, 0, 0, 0, 0, rule_hex, anchor)
    return gf
