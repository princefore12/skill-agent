"""Deck wrapper: a Presentation sized to the standard PowerPoint 16:9 slide."""

from pptx import Presentation
from pptx.util import Pt

from .units import SLIDE_PT_WIDTH, SLIDE_PT_HEIGHT
from .tokens import TOKENS


def strip_unused_layouts(prs):
    """Leave the delivered file with just the slide master (its theme + fonts)
    and the single **Blank** layout every slide is built on — drop the ~10 unused
    default layouts (Title Slide, Two Content, Comparison, …) so the PowerPoint
    Slide Master view is clean (client rule 2026-08-07). A layout that is in use
    by a slide is never removed (python-pptx guards it), so only the orphan
    templates go; the master and its fonts stay untouched.
    """
    for master in prs.slide_masters:
        for layout in list(master.slide_layouts):
            if not layout.used_by_slides:
                master.slide_layouts.remove(layout)


def clean_master(prs):
    """Strip the chrome the design system never uses (client rule 2026-08-07):

    * from every slide MASTER, remove the footer band — the date, footer, and
      slide-number placeholders — while keeping the title/body placeholders that
      carry the theme text styles and fonts;
    * from the Blank layout, remove EVERY placeholder so it is genuinely blank.

    Slides draw their own footer/pagination as plain text boxes (not the master
    placeholder), so nothing visible on a slide is lost.
    """
    from pptx.enum.shapes import PP_PLACEHOLDER
    band = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
    for master in prs.slide_masters:
        for ph in list(master.placeholders):
            if ph.placeholder_format.type in band:
                ph._element.getparent().remove(ph._element)
        for layout in master.slide_layouts:          # after strip: only Blank
            for ph in list(layout.placeholders):
                ph._element.getparent().remove(ph._element)


class Deck:
    """Thin wrapper over python-pptx Presentation.

    The slide is 960 x 540 pt (13.333 x 7.5 in) -- standard 16:9 widescreen.
    See units.py for why the px->pt factor is derived (0.5) rather than the
    inconsistent token value (0.75).
    """

    #: index of the blank slide layout in the default template
    _BLANK_LAYOUT = 6

    def __init__(self, tokens=TOKENS, grid_guides=True, brand_theme=True,
                 strip_layouts=True):
        self.tokens = tokens
        #: emit Prznt Perfect grid guides into viewProps.xml on save (client default: on)
        self.grid_guides = grid_guides
        #: overwrite the theme color scheme with the Prznt Perfect palette on
        #: save, so the PowerPoint color picker offers only design-system colors
        self.brand_theme = brand_theme
        #: on save, strip every unused slide layout (leaving just the master +
        #: the Blank layout) AND clean the chrome — the footer band off the master
        #: and every placeholder off the Blank layout (client rule 2026-08-07)
        self.strip_layouts = strip_layouts
        self.prs = Presentation()
        self.prs.slide_width = Pt(SLIDE_PT_WIDTH)
        self.prs.slide_height = Pt(SLIDE_PT_HEIGHT)

    def add_slide(self):
        layout = self.prs.slide_layouts[self._BLANK_LAYOUT]
        return self.prs.slides.add_slide(layout)

    def save(self, path):
        if self.brand_theme:
            from .theme import apply_brand_theme
            apply_brand_theme(self.prs, self.tokens)
        if self.grid_guides:
            from .guides import add_grid_guides
            add_grid_guides(self.prs)
        if self.strip_layouts:
            strip_unused_layouts(self.prs)
            clean_master(self.prs)
        self.prs.save(path)
        return path
